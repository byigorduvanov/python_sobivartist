from fastapi import Request, Depends, Form
from sqlalchemy.orm import Session
from starlette.responses import RedirectResponse
from starlette.status import HTTP_303_SEE_OTHER, HTTP_302_FOUND
from database.base import get_db
from main import app, templates
from tables.tables import Collective, Serie, Vavle, Material, Detail, Norm_year, Detail_change, Material_coming
from tables.tables_zero import Specification, Cost
import pandas as pd
from openpyxl import load_workbook


@app.get('/')
def home(request: Request):
    return templates.TemplateResponse('index.html', {'request': request})
@app.get('/addall')
def addall(request: Request, db_session: Session = Depends(get_db)):
    materials = db_session.query(Material).all()
    details = db_session.query(Detail).all()
    collectives = db_session.query(Collective).all()
    return templates.TemplateResponse('/add_all/index_add_all.html',
                                      {'request': request,
                                       'detail_list': details,
                                       'material_list': materials,}
                                      )
@app.get('/specification')
def specification(request: Request, db_session: Session = Depends(get_db)):
    materials = db_session.query(Material).all()
    details = db_session.query(Detail).all()
    collectives = db_session.query(Collective).all()   
    specifications = db_session.query(Specification).all()   
    return templates.TemplateResponse('index_specification.html',
                                      {'request': request,
                                       'detail_list': details,
                                       'material_list': materials,
                                       'collective_list':collectives,
                                       'specification_list':specifications}
                                      )
@app.get('/cost')
def cost(request: Request, db_session: Session = Depends(get_db)):
    materials = db_session.query(Material).all()
    details = db_session.query(Detail).all()
    collectives = db_session.query(Collective).all()
    series = db_session.query(Serie).all()
    vavles = db_session.query(Vavle).all()
    specifications = db_session.query(Specification).all()
    costs = db_session.query(Cost).all()    
    return templates.TemplateResponse('index_cost.html',
                                      {'request': request,
                                       'detail_list': details,
                                       'material_list': materials,
                                       'collective_list':collectives,
                                       'specification_list':specifications,
                                       'cost_list': costs}
                                      )
@app.get('/zminu')
def zminu(request: Request, db_session: Session = Depends(get_db)):
    return templates.TemplateResponse('/zminu/index_zminu.html',
                                      {'request': request}
                                      )

@app.get('/detail_all')
def detailall(request: Request, db_session: Session = Depends(get_db)):
    materials = db_session.query(Material).all()
    details = db_session.query(Detail).all()
    return templates.TemplateResponse('index_detail_all.html',
                                      {'request': request,
                                       'detail_list': details,
                                       'material_list': materials,}
                                      )


@app.get('/material_all')
def materialall(request: Request, db_session: Session = Depends(get_db)):
    materials = db_session.query(Material).all()
    details = db_session.query(Detail).all()
    return templates.TemplateResponse('index_material_all.html',
                                      {'request': request,
                                       'detail_list': details,
                                       'material_list': materials,}
                                      )                            


@app.get('/norm_year_all')
def norm_yearall(request: Request, db_session: Session = Depends(get_db)):
    materials = db_session.query(Material).all()
    details = db_session.query(Detail).all()
    norm_years = db_session.query(Norm_year).all()
    return templates.TemplateResponse('index_norm_year_all.html',
                                      {'request': request,
                                       'detail_list': details,
                                       'material_list': materials,
                                       'norm_year_list': norm_years,}
                                      )                 

@app.get('/revision')
def revision(request: Request, db_session: Session = Depends(get_db)):
    materials = db_session.query(Material).all()
    details = db_session.query(Detail).all()
    return templates.TemplateResponse('/revision/index_revision.html',
                                      {'request': request,
                                       'detail_list': details,
                                       'material_list': materials,}
                                      )
@app.post('/specification')
def specification(A:int = Form(...), db_session: Session = Depends(get_db)):
    for row in db_session.query(Detail, Collective).filter(Collective.detail_id == Detail.id).filter(Collective.collective_id == A).group_by(Collective.sequence_number):
        new_Specification = Specification(coding = row.Detail.coding, name = row.Detail.name, number = row.Collective.number)
        db_session.add(new_Specification)
        db_session.commit()
        print(new_Specification.coding, new_Specification.name, new_Specification.number)
       

    url = app.url_path_for('specification')
    return RedirectResponse(url=url, status_code=HTTP_303_SEE_OTHER)


@app.post('/dele_spec')
def dele_spec(db_session: Session = Depends(get_db)):
    new_specification = db_session.query(Specification).filter(Specification.id > 0)
    for i in new_specification:
        db_session.delete(i)
    db_session.commit()

    url = app.url_path_for('specification')
    return RedirectResponse(url=url, status_code=HTTP_302_FOUND)


@app.post('/cost')
def cost(A:int = Form(...), db_session: Session = Depends(get_db)):
   # new_collective_id = Collective(collective_id = A)
    for row in db_session.query(Detail, Collective, Material, Norm_year).filter(Collective.detail_id == Detail.id, Detail.material_id == Material.id, Norm_year.detail_id == Detail.id).filter(Collective.collective_id == A).group_by(Collective.sequence_number):
        rate_costs = round(((row.Material.weight /1000) * row.Detail.rate_costs), 4)
        rates_costs = rate_costs * row.Collective.number
  
        if row.Detail.price == 0:
            cost_materials = round((rates_costs * (row.Material.price /1000)), 2)                  
        else:
            cost_materials = row.Detail.price   
        norms_years = round(row.Norm_year.norm_years * row.Collective.number, 4)
        norms_sec = round(norms_years * 3600, 2)
        price_kg = round((row.Material.price /1000), 3)
   
       
        new_Cost = Cost(coding = row.Detail.coding, name = row.Detail.name, number = row.Collective.number, marka = row.Material.marka, profile = row.Material.profile, rate_costs = rate_costs, rates_costs = rates_costs, cost_materials = cost_materials, norm_years = row.Norm_year.norm_years, norms_years = norms_years, norms_sec = norms_sec, price_kg = price_kg)
        db_session.add(new_Cost)
        db_session.commit()

    url = app.url_path_for('cost')
    return RedirectResponse(url=url, status_code=HTTP_303_SEE_OTHER)


@app.post('/dele_cost')
def dele_cost( db_session: Session = Depends(get_db)):
    db_session.query(Cost).filter(Cost.id > 0).delete(synchronize_session="fetch")
    db_session.commit()
    
    url = app.url_path_for('cost')
    return RedirectResponse(url=url, status_code=HTTP_302_FOUND)


@app.post('/excel')
def excel(A:int = Form(...), db_session: Session = Depends(get_db)):
    #new_collective_id = Collective(collective_id = A)
    for row in db_session.query(Detail, Collective, Material, Norm_year).filter(Collective.detail_id == Detail.id, Detail.material_id == Material.id, Norm_year.detail_id == Detail.id).filter(Collective.collective_id == A).group_by(Collective.sequence_number):
        rate_costs = round(((row.Material.weight /1000) * row.Detail.rate_costs), 4)
        rates_costs = rate_costs * row.Collective.number
  
        if row.Detail.price == 0:
            cost_materials = round((rates_costs * (row.Material.price /1000)), 2)                  
        else:
            cost_materials = row.Detail.price   
        norms_years = round(row.Norm_year.norm_years * row.Collective.number, 4)
        norms_sec = round(norms_years * 3600, 2)
        price_kg = round((row.Material.price /1000), 3)
         
         
        #ПЕРЕЗАПИСУЄ ДАННІ 
        df = pd.DataFrame([row.Detail.coding, row.Detail.name, row.Collective.number, row.Material.marka, row.Material.profile, rate_costs, rates_costs, row.Norm_year.norm_years, cost_materials, norms_years, norms_sec, price_kg])
        df.to_excel('cost.xlsx')
        
        #ДОДАЄ ДАННІ
        wb = load_workbook('data1.xlsx')
        ws = wb['data']           
        ws.append([row.Detail.coding, row.Detail.name, row.Collective.number, row.Material.marka, row.Material.profile, rate_costs, rates_costs, row.Norm_year.norm_years, cost_materials, norms_years, norms_sec, price_kg])
        wb.save('data1.xlsx')
        wb.close() 

    url = app.url_path_for('cost')
    return RedirectResponse(url=url, status_code=HTTP_302_FOUND)
